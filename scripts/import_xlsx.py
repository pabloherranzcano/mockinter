#!/usr/bin/env python3
"""
Convierte Temario_mockinter.xlsx en un archivo SQL para cargar las preguntas en Supabase.

Uso:
    pip install openpyxl
    python scripts/import_xlsx.py Temario_mockinter.xlsx                       # genera supabase/seed.sql
    python scripts/import_xlsx.py Temario_mockinter.xlsx --overwrite           # el Excel pisa lo editado en el panel
    python scripts/import_xlsx.py Temario_mockinter.xlsx --overwrite --prune   # además borra temas y preguntas que no estén en el Excel

Por defecto las preguntas que ya existen en la base de datos NO se tocan, para no pisar
lo que hayas cambiado desde el panel de admin. Solo se rellenan las variantes en español vacías.

Las filas con Estado = «Descartada» se omiten. Activa = «No» se importa oculta.
El tema de cada pregunta sale de su ID: «ang-signals-M2» pertenece al tema «ang-signals».
"""
import sys
from pathlib import Path
from openpyxl import load_workbook

BLOCK_CODES = {
    "Frontend Angular": "ANG", "Frontend React": "REA", "Cloud / DevOps (AWS)": "CLD",
    "Backend Java/Spring": "BCK", "Frontend (Angular y React)": "FE",
    "Desarrollo (todos los puestos técnicos)": "DEV", "Común (todos los puestos)": "COM",
}
TYPES = {"Elegir respuesta": "mc", "Completar frase": "fill", "Ordenar": "order"}
STATUS = {"Pendiente de revisar": "pending", "Revisada": "reviewed", "Descartada": "discarded"}
ENGLISH_TOPICS = {"com-ingles"}


def q(v):
    """Literal SQL de texto (o NULL)."""
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "null"
    return "'" + str(v).replace("'", "''") + "'"


def arr(values):
    vals = [v for v in values if v not in (None, "")]
    if not vals:
        return "'{}'::text[]"
    return "array[" + ", ".join(q(v) for v in vals) + "]::text[]"


def main():
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    src = Path(sys.argv[1]); overwrite = "--overwrite" in sys.argv; prune = "--prune" in sys.argv
    ws = load_workbook(src, data_only=True)["Preguntas"]
    rows = list(ws.iter_rows(values_only=True))
    head = {h: i for i, h in enumerate(rows[0])}
    col = lambda r, name: r[head[name]] if name in head else None

    topics, questions, skipped, positions = {}, [], 0, {}
    for r in rows[1:]:
        qid = col(r, "ID")
        if not qid:
            continue
        if col(r, "Estado") == "Descartada":
            skipped += 1; continue
        block = col(r, "Bloque")
        code = BLOCK_CODES.get(block)
        if not code:
            print(f"Aviso: bloque desconocido en {qid}: {block}"); continue
        topic_id = str(qid).rsplit("-", 1)[0]
        if topic_id not in topics:
            positions[code] = positions.get(code, 0) + 1
            topics[topic_id] = (code, block, col(r, "Tema"), positions[code], topic_id in ENGLISH_TOPICS)
        t = TYPES.get(col(r, "Tipo"))
        if not t:
            print(f"Aviso: tipo desconocido en {qid}: {col(r, 'Tipo')}"); continue
        questions.append(dict(
            id=qid, topic_id=topic_id, level=col(r, "Nivel"), type=t,
            prompt=col(r, "Pregunta (ES)"), code=col(r, "Código"),
            correct=col(r, "Respuesta correcta"),
            wrong=[col(r, "Incorrecta 1"), col(r, "Incorrecta 2"), col(r, "Incorrecta 3")],
            explanation=col(r, "Explicación"), say_en=col(r, "Frase ejemplo (EN)"),
            say_es=col(r, "Traducción del ejemplo"),
            status=STATUS.get(col(r, "Estado"), "pending"), active=(col(r, "Activa") != "No"),
            es_prompt=col(r, "Variante ES: pregunta"), es_correct=col(r, "Variante ES: correcta"),
            es_wrong=[w.strip() for w in str(col(r, "Variante ES: incorrectas (separadas por |)") or "").split("|") if w.strip()],
            es_explanation=col(r, "Variante ES: explicación"),
        ))

    out = ["-- Generado por scripts/import_xlsx.py desde " + src.name, "begin;"]
    for tid, (code, block, name, pos, is_en) in topics.items():
        out.append(
            f"insert into public.topics (id, role, role_name, name, position, is_english) values "
            f"({q(tid)}, {q(code)}, {q(block)}, {q(name)}, {pos}, {'true' if is_en else 'false'}) "
            f"on conflict (id) do update set role = excluded.role, role_name = excluded.role_name, name = excluded.name, "
            f"position = excluded.position, is_english = excluded.is_english;")
    cols = ["id", "topic_id", "level", "type", "prompt", "code", "correct", "wrong", "explanation",
            "say_en", "say_es", "status", "active", "es_prompt", "es_correct", "es_wrong", "es_explanation"]
    conflict = ("do update set " + ", ".join(f"{c} = excluded.{c}" for c in cols[1:])) if overwrite else "do nothing"

    def val(d, c):
        if c == "wrong": return arr(d[c])
        if c == "es_wrong": return "null" if not d[c] else arr(d[c])
        if c == "active": return "true" if d[c] else "false"
        return q(d[c])

    for d in questions:
        out.append(f"insert into public.questions ({', '.join(cols)}) values ({', '.join(val(d, c) for c in cols)}) on conflict (id) {conflict};")
    if not overwrite:
        # Rellena las variantes en español vacías sin pisar lo escrito a mano
        for d in questions:
            sets = [f"{c} = coalesce({c}, {val(d, c)})" for c in ("es_prompt", "es_correct", "es_wrong", "es_explanation") if d[c]]
            if sets:
                out.append(f"update public.questions set {', '.join(sets)} where id = {q(d['id'])};")
    if prune:
        ids = ", ".join(q(d["id"]) for d in questions)
        tids = ", ".join(q(t) for t in topics)
        out.append(f"delete from public.questions where id not in ({ids});")
        out.append(f"delete from public.topics where id not in ({tids});")
    out.append("commit;")

    dest = Path(__file__).resolve().parent.parent / "supabase" / "seed.sql"
    dest.write_text("\n".join(out) + "\n", encoding="utf-8")
    flags = [f for f in ("--overwrite", "--prune") if f in sys.argv]
    print(f"{len(topics)} temas y {len(questions)} preguntas -> {dest} ({skipped} descartadas omitidas)"
          + (f" [{' '.join(flags)}]" if flags else ""))


if __name__ == "__main__":
    main()
